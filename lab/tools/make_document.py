"""
title: Make Document
author: agents-team
version: 0.1.0
description: Делает НАСТОЯЩИЙ файл — pptx, pdf, docx, xlsx, csv — и отдаёт его на скачивание. Нужен, когда просят «сгенерируй презентацию/PDF/таблицу», а не код для этого.
requirements: python-docx
"""
# ЗАЧЕМ. На демо человек попросил «сгенерируй pptx про пользу пива», потом «сгенерируй прям
# файл» — и получил ТРИ ответа подряд с питоновским кодом и финальным «я не могу отправить
# бинарный файл». Модель была права: сама по себе она умеет только текст. Файл должен делать
# инструмент на сервере.
#
# ПРАВИЛО ЭТОГО РЕПОЗИТОРИЯ (уже стоило нам трёх багов): модель НИКОГДА не копирует крупный
# payload в ответ — она упирается в max_tokens и обрывается посреди. Поэтому здесь всё то же,
# что в diagram_tool/project_gen: файл пишется через Files API на сервере, в чат уходит
# карточка файла и короткая ссылка.
#
# ФОРМАТЫ И ЧЕМ ДЕЛАЕМ (проверено, что лежит в образе):
#   pptx  — python-pptx   (есть)
#   xlsx  — openpyxl      (есть)
#   pdf   — fpdf2         (есть; кириллица только со своим TTF — берём DejaVu из образа)
#   docx  — python-docx   (НЕТ в образе, ставится из фронтматтера `requirements`)
#   csv/md — просто текст
#
# Вход у всех один — Markdown. Модели он даётся заметно надёжнее, чем вложенный JSON со
# слайдами: меньше поводов сломать tool-call на экранировании.
import csv as _csv
import io
import re

_FONTS = [
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/app/backend/open_webui/static/fonts/NotoSans-Regular.ttf",
]
_FONTS_BOLD = [
    "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
    "/app/backend/open_webui/static/fonts/NotoSans-Bold.ttf",
]
_KINDS = ("pptx", "pdf", "docx", "xlsx", "csv", "md")


def _unfence(text: str) -> str:
    """Модель часто заворачивает содержимое в ```markdown … ``` — снимаем обёртку."""
    t = (text or "").strip()
    m = re.match(r"^```[a-zA-Z]*\n(.*)\n```$", t, re.S)
    return m.group(1).strip() if m else t


def _slug(name: str, default: str) -> str:
    s = re.sub(r"[^\w \-.]+", "", (name or "").strip()).strip()
    s = re.sub(r"\s+", "_", s)[:60]
    return s or default


# Модель сплошь и рядом игнорирует «каждый ## начинает слайд» и пишет по-своему:
# «СЛАЙД 1. Название», «Slide 3:», «1. Название». Проверено на живом запросе про 25
# слайдов: `##` не было ни одного, и вся колода схлопывалась в два слайда. Поэтому
# распознаём и такие маркеры, иначе формально работающая тулза даёт мусор на выходе.
_SLIDE_MARK = re.compile(
    r"^(?:слайд|slide)\s*№?\s*\d+\s*[.:)-]?\s*(.*)$", re.I)
_NUM_HEAD = re.compile(r"^\d{1,2}\s*[.)]\s+(\S.*)$")


def _blocks(md: str):
    """Markdown → плоский список блоков (заголовок / пункт / абзац / код)."""
    out, in_code, buf = [], False, []
    for raw in md.splitlines():
        line = raw.rstrip()
        if line.strip().startswith("```"):
            if in_code:
                out.append(("code", "\n".join(buf)))
                buf = []
            in_code = not in_code
            continue
        if in_code:
            buf.append(raw)
            continue
        s = line.strip()
        if not s:
            continue
        h = re.match(r"^(#{1,6})\s+(.*)$", s)
        if h:
            out.append((f"h{len(h.group(1))}", h.group(2).strip()))
            continue
        m = _SLIDE_MARK.match(s)
        if m:                                   # «СЛАЙД 4. Название» = заголовок слайда
            out.append(("h2", (m.group(1) or f"Слайд").strip(" .:-")))
            continue
        b = re.match(r"^[-*+]\s+(.*)$", s)
        if b:
            out.append(("li", b.group(1).strip()))
            continue
        n = _NUM_HEAD.match(s)
        if n:
            # «1. Текст» — заголовок, только если это короткая строка без точки в конце;
            # иначе это обычный нумерованный пункт списка.
            head = n.group(1).strip()
            if len(head) <= 70 and not head.endswith((".", "!", "?", ";")):
                out.append(("h2", head))
            else:
                out.append(("li", head))
            continue
        out.append(("p", s))
    if in_code and buf:
        out.append(("code", "\n".join(buf)))
    return out


def _plain(s: str) -> str:
    """Снять инлайновую разметку: в pptx/pdf «**жирный**» звёздочками выглядит мусором."""
    s = re.sub(r"\*\*(.+?)\*\*", r"\1", s)
    s = re.sub(r"(?<!\*)\*(?!\s)(.+?)(?<!\s)\*", r"\1", s)
    s = re.sub(r"`(.+?)`", r"\1", s)
    s = re.sub(r"\[(.+?)\]\((.+?)\)", r"\1", s)
    return s


# --------------------------------------------------------------------------- pptx
def _build_pptx(title: str, md: str) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)   # 16:9

    def title_slide(t, sub):
        s = prs.slides.add_slide(prs.slide_layouts[0])
        s.shapes.title.text = t
        if len(s.placeholders) > 1:
            s.placeholders[1].text = sub

    def body_slide(t, items):
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = t
        box = s.shapes.add_textbox(Inches(0.8), Inches(1.7), Inches(11.7), Inches(5.2))
        tf = box.text_frame
        tf.word_wrap = True
        for i, (kind, txt) in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = ("• " + txt) if kind == "li" else txt
            p.font.size = Pt(20 if kind == "li" else 18)
            p.space_after = Pt(10)

    # Сначала режем markdown на секции «заголовок + его содержимое», и только потом
    # решаем, что из этого титульник. Иначе подзаголовок сразу после `#` становился
    # отдельным пустым слайдом (так и вышло на первом же боевом запросе).
    sections, head, items, lead = [], None, [], []
    for kind, txt in _blocks(md):
        txt = _plain(txt)
        if kind.startswith("h"):
            if head is not None:
                sections.append((head, items))
            head, items = txt, []
        elif head is None:
            lead.append(("li" if kind == "li" else "p", txt))
        else:
            items.append(("li" if kind == "li" else "p", txt))
    if head is not None:
        sections.append((head, items))

    if not sections:
        title_slide(title or "Презентация", "")
        if lead:
            body_slide("Содержание", lead)
    else:
        head, items = sections[0]
        rest = sections[1:]
        # Пустой заголовок сразу за титульным — это подзаголовок, а не отдельный слайд.
        sub = ""
        if not items and rest and not rest[0][1]:
            sub, rest = rest[0][0], rest[1:]
        title_slide(title or head, sub)
        if items:                      # у первой секции был текст — он не должен пропасть
            body_slide(head, items)
        for h, its in rest:
            body_slide(h, its)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------- pdf
def _build_pdf(title: str, md: str) -> bytes:
    import os
    from fpdf import FPDF

    reg = next((p for p in _FONTS if os.path.exists(p)), None)
    bold = next((p for p in _FONTS_BOLD if os.path.exists(p)), None)
    if not reg:
        raise RuntimeError("в контейнере нет TTF с кириллицей")

    pdf = FPDF(format="A4")
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.add_font("body", "", reg)
    pdf.add_font("body", "B", bold or reg)
    pdf.add_page()
    w = pdf.w - pdf.l_margin - pdf.r_margin

    if title:
        pdf.set_font("body", "B", 20)
        pdf.multi_cell(w, 10, _plain(title))
        pdf.ln(3)
    for kind, txt in _blocks(md):
        if kind == "code":
            pdf.set_font("body", "", 9)
            pdf.set_fill_color(244, 244, 244)
            pdf.multi_cell(w, 5, txt, fill=True)
            pdf.ln(2)
            continue
        txt = _plain(txt)
        if kind.startswith("h"):
            lvl = int(kind[1])
            pdf.ln(3)
            pdf.set_font("body", "B", max(11, 19 - 2 * lvl))
            pdf.multi_cell(w, 8, txt)
            pdf.ln(1)
        elif kind == "li":
            pdf.set_font("body", "", 11)
            pdf.multi_cell(w, 6, "•  " + txt)
        else:
            pdf.set_font("body", "", 11)
            pdf.multi_cell(w, 6, txt)
            pdf.ln(1)
    out = pdf.output()
    return bytes(out) if not isinstance(out, (bytes, bytearray)) else bytes(out)


# --------------------------------------------------------------------------- docx
def _build_docx(title: str, md: str) -> bytes:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    if title:
        doc.add_heading(_plain(title), level=0)
    for kind, txt in _blocks(md):
        if kind == "code":
            p = doc.add_paragraph()
            r = p.add_run(txt)
            r.font.name = "Consolas"
            r.font.size = Pt(9)
        elif kind.startswith("h"):
            doc.add_heading(_plain(txt), level=min(int(kind[1]), 4))
        elif kind == "li":
            doc.add_paragraph(_plain(txt), style="List Bullet")
        else:
            doc.add_paragraph(_plain(txt))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------- xlsx
def _rows(text: str):
    text = _unfence(text)
    # Markdown-таблица | a | b | — модель отдаёт её охотнее, чем CSV
    if "|" in text and re.search(r"^\s*\|", text, re.M):
        rows = []
        for line in text.splitlines():
            s = line.strip()
            if not s.startswith("|"):
                continue
            if re.match(r"^\|[\s:\-|]+\|$", s):     # разделитель шапки
                continue
            rows.append([c.strip() for c in s.strip("|").split("|")])
        if rows:
            return rows
    delim = "\t" if text.count("\t") > text.count(",") else ","
    return [r for r in _csv.reader(io.StringIO(text), delimiter=delim) if r]


def _build_xlsx(title: str, text: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = (_plain(title) or "Лист")[:31]
    rows = _rows(text)
    for r in rows:
        ws.append(r)
    if rows:
        for c in ws[1]:
            c.font = Font(bold=True)
        for i, _ in enumerate(rows[0], start=1):
            width = max((len(str(r[i - 1])) for r in rows if len(r) >= i), default=10)
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(60, width + 2)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _looks_truncated(text: str):
    """
    Оборвался ли текст на полуслове.

    Тулза не видит finish_reason, а модель на длинном документе упирается в
    max_tokens ПРЯМО ВНУТРИ аргумента tool-call. Парсер серва при этом закрывает
    JSON, вызов проходит как валидный, и раньше мы молча собирали обрезанный файл.
    Проверено вживую: на max_tokens=4096 запрос «25 слайдов» дал finish=length и
    текст, оборванный посреди предложения.
    """
    t = (text or "").rstrip()
    if not t:
        return None
    if t.count("```") % 2 == 1:
        return "не закрыт блок кода"
    # Повтор проверяем раньше остального: зациклившийся текст тоже кончается на букве.
    lines = [l.strip() for l in t.splitlines() if l.strip()]
    if len(lines) >= 6 and len(set(lines[-5:])) == 1:
        return "модель зациклилась на повторе"
    if t[-1] in ",-—:;(«“":
        return "текст обрывается на связке"
    # 🚨 «кончается на букву» НЕЛЬЗЯ считать обрывом: буллиты слайдов сплошь и рядом
    # без точки в конце. Проверено на семи реальных документах — шесть из них так и
    # заканчиваются, и грубая проверка ругалась на исправные файлы.
    # Настоящая примета обрыва по лимиту — оборванное СЛОВО: генерация встаёт на
    # произвольном токене (боевой пример: «…английский инженер Джордж К»).
    last = re.split(r"[\s]+", t)[-1].strip(".,!?;:)»\"'")
    if last and last.isalpha() and len(last) <= 2:
        # Регистр разводит обрубок и настоящий предлог: «Джордж К» (обрыв фамилии)
        # против «ведёт к» (законный конец строки). Односимвольный обрубок обычно
        # сохраняет заглавную букву исходного слова.
        if last[0].isupper() or last.lower() not in _SHORT_WORDS:
            return "последнее слово оборвано"
    return None


# Настоящие короткие слова, чтобы не принимать их за обрубок.
_SHORT_WORDS = {
    "в", "и", "к", "с", "о", "у", "я", "а", "на", "за", "до", "по", "из", "не", "ни",
    "но", "то", "же", "ли", "бы", "их", "им", "ею", "ей", "ом", "их",
    "a", "i", "an", "to", "of", "in", "is", "it", "we", "he", "as", "at", "by", "or",
    "on", "if", "so", "up", "no", "do", "me", "my", "us",
}


_MIME = {
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
    "csv": "text/csv",
    "md": "text/markdown",
}


class Tools:
    def __init__(self):
        self.citation = False

    async def make_document(
        self,
        kind: str,
        title: str,
        content: str,
        append_to: str = "",
        __request__=None,
        __user__=None,
        __event_emitter__=None,
        __chat_id__: str = "",
    ) -> str:
        """
        Создать настоящий файл документа и отдать пользователю на скачивание. Вызывай, когда просят сгенерировать презентацию, PDF, документ Word, таблицу Excel или CSV — сам файл, а не код для его создания.
        :param kind: Формат файла: pptx (презентация), pdf, docx (Word), xlsx (Excel), csv, md.
        :param title: Заголовок документа и основа имени файла, например «Польза пива».
        :param content: Содержимое в Markdown. Для pptx каждый заголовок (## Название) начинает новый слайд, пункты списка становятся буллитами. Для xlsx и csv передавай таблицу: markdown-таблицу или CSV со строкой заголовков.
        :param append_to: Необязательно. Идентификатор файла из прошлого вызова, чтобы ДОПИСАТЬ к нему content, а не создавать документ заново. Так делают длинные документы по частям и добавляют разделы к уже готовому файлу.
        """
        kind = (kind or "").strip().lower().lstrip(".")
        if kind in ("powerpoint", "ppt", "презентация"):
            kind = "pptx"
        if kind in ("word", "doc"):
            kind = "docx"
        if kind in ("excel", "xls", "таблица"):
            kind = "xlsx"
        if kind not in _KINDS:
            return f"Не умею формат «{kind}». Доступны: {', '.join(_KINDS)}."

        content = _unfence(content)
        if not content.strip():
            return "Пустое содержимое: нечего класть в файл."

        async def st(desc, done=False):
            if __event_emitter__:
                await __event_emitter__({"type": "status",
                                         "data": {"description": desc, "done": done}})

        # ДОПИСЫВАНИЕ. Исходный markdown храним в data-файла, поэтому продолжение
        # не требует от модели заново выдавать весь документ — а именно это и упирается
        # в max_tokens на длинных доках.
        prefix = ""
        if append_to:
            try:
                from open_webui.models.files import Files as _F
                prev = await _F.get_file_by_id(append_to.strip())
                prefix = ((prev.data or {}).get("content") or "") if prev else ""
            except Exception:
                prefix = ""
            if not prefix:
                return (f"Не нашёл исходник документа {append_to} и не могу дописать. "
                        "Собери файл заново одним вызовом без append_to.")
            content = prefix.rstrip() + "\n\n" + content.lstrip()

        cut = _looks_truncated(content)

        await st(f"собираю {kind}…")
        try:
            if kind == "pptx":
                data = _build_pptx(title, content)
            elif kind == "pdf":
                data = _build_pdf(title, content)
            elif kind == "docx":
                data = _build_docx(title, content)
            elif kind == "xlsx":
                data = _build_xlsx(title, content)
            else:
                data = content.encode("utf-8")
        except ImportError as e:
            await st("не хватает библиотеки", True)
            return (f"Не могу собрать {kind}: не установлена библиотека ({e}). "
                    f"Предложи пользователю другой формат — pdf, pptx и xlsx работают.")
        except Exception as e:
            await st("ошибка сборки", True)
            return f"Не удалось собрать {kind}: {type(e).__name__}: {e}"

        name = f"{_slug(title, 'document')}.{kind}"
        if not __request__ or not __user__:
            return (f"Файл {name} собран ({len(data)} байт), но сохранить некуда: "
                    "нет request/user (не UI-контур).")
        try:
            from fastapi import UploadFile
            from open_webui.models.users import Users
            from open_webui.models.files import Files
            from open_webui.routers.files import upload_file_handler

            user = await Users.get_user_by_id(__user__["id"])
            up = UploadFile(file=io.BytesIO(data), filename=name,
                            headers={"content-type": _MIME.get(kind, "application/octet-stream")})
            # process=False: индексировать документ в RAG не надо, он и так у пользователя.
            item = await upload_file_handler(__request__, file=up,
                                             metadata={"chat_id": __chat_id__, "lab_document": True},
                                             process=False, process_in_background=False, user=user)
            # Исходный markdown кладём ВСЕГДА: это то, из чего потом делается append.
            await Files.update_file_data_by_id(item.id, {"content": content})
        except Exception as e:
            await st("ошибка сохранения", True)
            return f"Файл собрался, но не сохранился: {type(e).__name__}: {e}"

        url = f"/api/v1/files/{item.id}/content/{name}"
        if __event_emitter__:
            await __event_emitter__({"type": "files", "data": {"files": [
                {"type": "file", "id": item.id, "name": name, "size": len(data),
                 "url": f"/api/v1/files/{item.id}"}]}})
            await st(f"готово: {name} ({len(data) // 1024 or 1} КБ)", True)

        kb = len(data) // 1024 or 1
        warn = ""
        if cut:
            # Не молчим: обрезанный документ выглядит целым, и пользователь узнает об этом
            # последним. Заодно подсказываем модели готовый способ дописать хвост.
            warn = (f"\n\n⚠ ВНИМАНИЕ: содержимое похоже на оборванное ({cut}) — скорее всего "
                    f"ты упёрся в лимит токенов. Скажи об этом пользователю и допиши остаток "
                    f"вторым вызовом make_document с append_to=\"{item.id}\".")
        return (f"Файл готов и уже прикреплён к сообщению: **{name}** ({kb} КБ). "
                f"id={item.id}\n\n[Скачать {name}]({url}){warn}\n\n"
                f"Ответь пользователю ОДНОЙ-ДВУМЯ строками: файл готов, вот ссылка. "
                f"НЕ пересказывай содержимое файла и НЕ вставляй код. Чтобы дописать разделы "
                f"в этот же документ, вызови make_document с append_to=\"{item.id}\".")
